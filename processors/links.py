import io
import re

import docx
import pymupdf

from processors.text import DOCX, PDF

URL_PATTERN = re.compile(r"https?://[A-Za-z0-9._~:/?#\[\]@!$&'()*+,;=%-]*[A-Za-z0-9/_#=-]")
GITHUB_HOSTS = ("github.com", "gist.github.com", "raw.githubusercontent.com")
TRAILING = ".,;:)]}'\"`>"
CONTEXT_BEFORE = 400
CONTEXT_AFTER = 160


def tidy(url):
    url = url.strip().rstrip(TRAILING)
    return url[:-1] if url.endswith("/") else url


def is_github(url):
    host = url.split("//", 1)[-1].split("/", 1)[0].split("@")[-1].lower()
    return host in GITHUB_HOSTS


def rejoin(text):
    # pdf extraction breaks a long url at the line end, the hyphen before the break is part of the url
    return re.sub(r"(?<=[A-Za-z0-9._~:/?#@!$&'()*+,;=%-])\n[ \t]*(?=[A-Za-z0-9._~/?#=%-])", "", text)


def tighten(text):
    return re.sub(r"\s+", " ", text).strip()


def window(text, index, length):
    return tighten(text[max(0, index - CONTEXT_BEFORE):index + length + CONTEXT_AFTER])


def from_pdf(data):
    found = []
    with pymupdf.open(stream=data, filetype="pdf") as pdf:
        for page in pdf:
            text = page.get_text()
            for link in page.get_links():
                if link.get("kind") != pymupdf.LINK_URI or not link.get("uri"):
                    continue
                # the words the link is drawn over say what it points at, the url alone does not
                anchor = tighten(page.get_textbox(link["from"]))
                index = text.find(anchor.split(" ")[0]) if anchor else -1
                context = window(text, index, len(anchor)) if index >= 0 else ""
                found.append((tidy(link["uri"]), anchor, context))
    return found


def from_docx(data):
    document = docx.Document(io.BytesIO(data))
    targets = {
        rel_id: rel.target_ref
        for rel_id, rel in document.part.rels.items()
        if rel.reltype.endswith("/hyperlink") and rel.is_external and rel.target_ref.startswith("http")
    }
    paragraphs = [p.text for p in document.paragraphs]
    found = []
    for position, paragraph in enumerate(document.paragraphs):
        xml = paragraph._p.xml
        for rel_id, target in targets.items():
            if f'r:id="{rel_id}"' not in xml:
                continue
            before = paragraphs[position - 1] if position else ""
            found.append((tidy(target), tighten(paragraph.text), tighten(f"{before} {paragraph.text}")))
    return found


def from_text(text):
    joined = rejoin(text)
    return [(tidy(m.group()), "", window(joined, m.start(), len(m.group()))) for m in URL_PATTERN.finditer(joined)]


def from_message(message_text):
    # slack wraps urls as <url> or <url|label>
    text = message_text or ""
    found = [(tidy(u), label.strip(" |"), tighten(text))
             for u, label in re.findall(r"<(https?://[^>|]+)(\|[^>]*)?>", text)]
    return found + [(u, a, tighten(text)) for u, a, _ in from_text(re.sub(r"<[^>]*>", " ", text))]


def candidates(data, mimetype, text, message_text):
    harvested = []
    embedded = from_pdf(data) if mimetype == PDF else from_docx(data) if mimetype == DOCX else []
    harvested += [(u, a, c, "document link") for u, a, c in embedded]
    harvested += [(u, a, c, "document text") for u, a, c in from_text(text)]
    harvested += [(u, a, c, "slack message") for u, a, c in from_message(message_text)]
    found = {}
    for url, anchor, context, source in harvested:
        if not is_github(url):
            continue
        entry = found.setdefault(url, {"anchor": "", "context": "", "sources": []})
        if source not in entry["sources"]:
            entry["sources"].append(source)
        if anchor and not entry["anchor"]:
            entry["anchor"] = anchor
        if context and len(context) > len(entry["context"]):
            entry["context"] = context
    return found


def describe(found):
    lines = []
    for url, entry in found.items():
        lines.append(f"- {url}")
        lines.append(f"    found in: {', '.join(entry['sources'])}")
        if entry["anchor"]:
            lines.append(f"    the link is attached to the words: {entry['anchor']!r}")
        if entry["context"]:
            lines.append(f"    the passage it appears in: \"{entry['context']}\"")
    return "\n".join(lines)
